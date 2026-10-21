from pathlib import Path
from html import escape
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
P=Path(__file__).resolve().parents[1];OUT=P/'production';OUT.mkdir(exist_ok=True)
lines=(P/'manuscript/MASTER_DRAFT.md').read_text(encoding='utf8').splitlines()
blocks=[];para=[];table=[]
def flush():
 if para:blocks.append(('p',' '.join(para)));para.clear()
def tabflush():
 if table:
  rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in table if not re.match(r'^\|?\s*:?-{2,}',l.strip())];blocks.append(('table',rows));table.clear()
for line in lines+['']:
 t=line.strip()
 if t.startswith('|'):
  flush();table.append(t);continue
 tabflush()
 if not t:flush();continue
 if t=='<!-- pagebreak -->':flush();blocks.append(('pagebreak',None));continue
 m=re.match(r'^(#{1,6})\s+(.*)',t)
 if m:flush();blocks.append(('h'+str(len(m[1])),m[2]));continue
 m=re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$',t)
 if m:flush();blocks.append(('image',(m[1],m[2])));continue
 if t.startswith(('```','---')):flush();continue
 if t.startswith('>'):flush();blocks.append(('quote',t.lstrip('> ')));continue
 if re.match(r'^(?:[-*]|\d+[.)])\s+',t):flush();blocks.append(('li',t));continue
 para.append(t)

# Replace atlas blocks with physical landscape plates; retain full teaching notes afterwards.
plate_manifest=json.loads((P/'circuits/print_manifest.json').read_text(encoding='utf8'))
plates={d['id']:d for d in plate_manifest['diagrams']}
newblocks=[];notes=[];i=0;in_atlas=False
while i<len(blocks):
 kind,value=blocks[i]
 if kind=='h2' and value.startswith('ภาคหนึ่ง'):
  in_atlas=True;i+=1
  if i<len(blocks) and blocks[i][0]=='p':i+=1
  continue
 match=re.match(r'^ภาพ (\d+)',value) if kind=='h3' else None
 if in_atlas and match:
  key=match[1].zfill(2);d=plates[key];newblocks.append(('plate',d));notes.append(d);i+=1
  while i<len(blocks) and blocks[i][0]!='h2' and not (blocks[i][0]=='h3' and re.match(r'^ภาพ \d+',blocks[i][1])):i+=1
  continue
 if in_atlas and notes:
  newblocks.append(('h2','คำอธิบายประกอบแผ่นวงจร'))
  for d in notes:newblocks.extend([('h3','ภาพ '+d['id']+' '+d['title_th']),('p',d['teaching_note_th'])])
  in_atlas=False
 newblocks.append(blocks[i]);i+=1
blocks=newblocks

def inline(s):
 s=escape(s);s=re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'<a href="\2">\1</a>',s)
 s=re.sub(r'\[([^\]]+)\]\((?!https?://)[^)]+\)',r'\1',s)
 s=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',s);return re.sub(r'`([^`]+)`',r'<code>\1</code>',s)

doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(1.65);sec.left_margin=sec.right_margin=Cm(1.8)
for name,size in [('Normal',11.5),('Title',25),('Heading 1',18),('Heading 2',15),('Heading 3',13),('Caption',9)]:
 st=doc.styles[name];st.font.name='Leelawadee UI';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=name.startswith(('Title','Heading'))
 st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.18
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Leelawadee UI')
for el in list(doc.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
def clean(s):return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m:m[1]+' <'+m[2]+'>' if m[2].startswith(('http:','https:')) else m[1],s).replace('**','').replace('`','')
html=[];imcount=0;landscape=False

def page_orientation(wide):
 global landscape
 if wide==landscape:
  if wide:doc.add_page_break()
  return False
 section=doc.add_section(WD_SECTION_START.NEW_PAGE)
 section.orientation=WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
 section.page_width=Cm(29.7 if wide else 21);section.page_height=Cm(21 if wide else 29.7)
 section.top_margin=section.bottom_margin=Cm(1.2 if wide else 1.65)
 section.left_margin=section.right_margin=Cm(1.2 if wide else 1.8)
 landscape=wide
 return True

def wordlink(paragraph,label,url):
 link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
 run=OxmlElement('w:r');rpr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'18');rpr.append(size);run.append(rpr);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);paragraph._p.append(link)

for kind,value in blocks:
 if kind=='plate':
  if imcount==0:
   doc.add_page_break()
   tp=doc.add_paragraph('สารบัญ Word','Title')
   doc.add_paragraph('เลขหน้าของ Word อาจต่างจาก PDF ให้คลิกขวาที่สารบัญแล้วเลือกอัปเดตเขตข้อมูลหลังแก้ไขต้นฉบับ')
   tp=doc.add_paragraph();tf=OxmlElement('w:fldSimple');tf.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
   tr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='อัปเดตสารบัญใน Microsoft Word';tr.append(tt);tf.append(tr);tp._p.append(tf)
  d=value;page_orientation(True);path=P/'circuits'/d['png'];im=Image.open(path);w,h=im.size;ww=min(26.9,15*w/h)
  head=doc.add_paragraph('ภาพ '+d['id']+' '+d['title_th'],'Heading 3');head.paragraph_format.keep_with_next=True
  pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(2);pp.add_run().add_picture(str(path),width=Cm(ww))
  scope='ตัวอย่าง DR201E Hilux TASA Argentina เท่านั้น' if d['id'] in ['11','12'] else 'ผังหลักการ ไม่ใช่ผัง OEM สากล'
  pp=doc.add_paragraph(d['caption_th']+' • '+scope,'Caption');pp.paragraph_format.space_after=Pt(2)
  pp=doc.add_paragraph('แหล่งอ้างอิง ', 'Caption')
  for n,url in enumerate(d['sources'],1):wordlink(pp,'['+str(n)+'] ',url)
  if not d['sources']:pp.add_run('แบบจำลองวงจรที่ผู้เขียนวาดขึ้น')
  links=' • '.join('<a href="'+escape(u)+'">แหล่ง '+str(j+1)+'</a>' for j,u in enumerate(d['sources'])) or 'แบบจำลองวงจรที่ผู้เขียนวาดขึ้น'
  html.append('<section class="plate"><h3>ภาพ '+d['id']+' '+escape(d['title_th'])+'</h3><img src="'+path.as_uri()+'"><p class="plate-caption">'+escape(d['caption_th'])+' • '+scope+'</p><p class="plate-source">'+links+'</p></section>');imcount+=1;continue
 switched=page_orientation(False) if landscape else False
 if kind=='pagebreak':
  doc.add_page_break();html.append('<div class="pagebreak"></div>');continue
 if kind.startswith('h'):
  level=int(kind[1:]);style='Title' if level==1 else 'Heading '+str(min(level-1,3));p=doc.add_paragraph(clean(value),style);p.paragraph_format.keep_with_next=True
  if level==2 and len(html)>2 and not switched:p.paragraph_format.page_break_before=True
  if re.match(r'^แบบ0[1-6]',value) or value.startswith('บทเรียน 5'):p.paragraph_format.page_break_before=True
  html.append(f'<h{min(level,4)}'+(' class="field-start"' if re.match(r'^แบบ0[1-6]',value) or value.startswith('บทเรียน 5') else '')+'>'+inline(value)+f'</h{min(level,4)}>')
 elif kind=='image':
  alt,relative=value;path=(P/'manuscript'/relative).resolve();im=Image.open(path);w,h=im.size;ww=min(17.4,12*w/h)
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Cm(ww));html.append(f'<figure><img src="{path.as_uri()}" alt="{escape(alt)}"></figure>');imcount+=1
 elif kind=='table':
  if not value:continue
  n=max(map(len,value))
  if n>4:
   headers=value[0]
   for row in value[1:]:
    if not row:continue
    cardtitle=(headers[0]+' '+row[0]) if not row[0].strip('_ .') else row[0]
    h=doc.add_paragraph(clean(cardtitle),'Heading 3');h.paragraph_format.keep_with_next=True
    html.append('<article class="repair-card"><h4>'+inline(cardtitle)+'</h4>')
    for index,txt in enumerate(row[1:],1):
     label=headers[index] if index<len(headers) else 'ข้อมูล'
     pp=doc.add_paragraph();pp.add_run(clean(label)+' ').bold=True;pp.add_run(clean(txt));pp.paragraph_format.keep_together=True
     html.append('<p><b>'+inline(label)+'</b> '+inline(txt)+'</p>')
    html.append('</article>')
   continue
  tab=doc.add_table(rows=0,cols=n);tab.style='Table Grid' 
  for ri,row in enumerate(value):
   cells=tab.add_row().cells
   for ci,txt in enumerate(row):
    cells[ci].text=clean(txt)
    for p in cells[ci].paragraphs:
     for r in p.runs:r.font.size=Pt(9.5);r.bold=ri==0
  # Repeat header rows in both rendering paths.
  if tab.rows:
   hp=tab.rows[0]._tr.get_or_add_trPr();flag=OxmlElement('w:tblHeader');hp.append(flag)
  header='<thead><tr>'+''.join('<th>'+inline(c)+'</th>' for c in value[0])+'</tr></thead>'
  bodyrows='<tbody>'+''.join('<tr>'+''.join('<td>'+inline(c)+'</td>' for c in row)+'</tr>' for row in value[1:])+'</tbody>'
  sourceclass=' class="source-table"' if value[0][0].strip().lower() in ['id','รหัส','source id'] else ''
  fieldtable=n==2 and 'บันทึก' in value[0][1]
  if fieldtable:
   tab.autofit=False
   for rr in tab.rows:
    rr.cells[0].width=Cm(6);rr.cells[1].width=Cm(11.4)
  colspec='<colgroup><col style="width:35%"><col style="width:65%"></colgroup>' if fieldtable else ''
  html.append('<table'+sourceclass+(' style="table-layout:fixed"' if fieldtable else '')+'>'+colspec+header+bodyrows+'</table>')
 else:
  text=clean(value);p=doc.add_paragraph(text)
  if kind=='quote':
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFF2CC');p._p.get_or_add_pPr().append(shade)
  p.paragraph_format.keep_together=True
  tag='blockquote' if kind=='quote' else 'p';html.append(f'<{tag}'+(' class="list"' if kind=='li' else '')+'>'+inline(value)+f'</{tag}>')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('ต้นฉบับแก้ไข P0093 • ');run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='P0093 Toyota 1KD และ 2KD วงจรและการวินิจฉัย';doc.core_properties.author='กองบรรณาธิการโครงการ';doc.save(OUT/'P0093_Prototype.docx')
css='''@page diagram{size:A4 landscape;margin:12mm}@page{size:A4;margin:17mm 18mm 19mm}*{box-sizing:border-box}body{font-family:"Leelawadee UI",Tahoma,sans-serif;color:#000;font-size:11.5pt;line-height:1.36;margin:0}h1{font-size:25pt;margin:0 0 8mm;color:#000}h2{font-size:18pt;break-before:page;margin:0 0 6mm;color:#000}h3{font-size:14pt;margin:6mm 0 3mm;color:#000}h4{font-size:12.5pt;margin:5mm 0 2mm;color:#000}h1,h2,h3,h4{break-after:avoid}h3.figure-title{break-before:page}h2.subtitle{break-before:auto;font-size:17pt}h2+p+h3.figure-title{break-before:auto}p{margin:0 0 3mm;orphans:3;widows:3;overflow-wrap:anywhere}p.list{margin-left:4mm}figure{margin:3mm 0 5mm;break-inside:avoid}img{display:block;width:100%;max-height:122mm;object-fit:contain}table{width:100%;border-collapse:collapse;font-size:9.5pt;line-height:1.3;margin:4mm 0}td,th{border:1px solid #bbb;padding:2mm;overflow-wrap:anywhere;vertical-align:top}th{background:#f1f3f4}tr{break-inside:avoid}blockquote{margin:3mm 0;padding:3mm;background:#fff2cc}.repair-card{break-inside:avoid;margin:5mm 0;padding:3mm;border:1px solid #ddd}a{color:#155889;overflow-wrap:anywhere}code{font-family:inherit}'''
css+=''' .plate{page:diagram;width:273mm;break-before:page;break-after:page}.plate h3{margin:0 0 2mm;font-size:14pt}.plate img{width:269mm;height:150mm;max-height:none;object-fit:contain;margin:0 auto 1mm}.plate p{margin:0 0 1mm}.plate-caption{font-size:10pt;line-height:1.22}.plate-source{font-size:9pt;line-height:1.2}.pagebreak,.field-start{break-before:page}thead{display:table-header-group;break-after:avoid}tfoot{display:table-footer-group}.source-table td:first-child,.source-table th:first-child{white-space:nowrap;min-width:12mm} '''
body='\n'.join(html);body=body.replace('<h2>','<h2 class="subtitle">',1);body=body
(OUT/'prototype.html').write_text('<!doctype html><html lang="th"><meta charset="utf-8"><title>P0093 ต้นแบบ</title><style>'+css+'</style><body>'+body+'</body></html>',encoding='utf8')
(OUT/'build_manifest.json').write_text(json.dumps({'blocks':len(blocks),'figures':imcount,'status':'prototype not final'},ensure_ascii=False,indent=2),encoding='utf8')
print('Prototype DOCX and HTML:',len(blocks),'blocks',imcount,'figures')
